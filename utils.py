import logging
import re
import sys
from pathlib import Path
from typing import List, Dict, Any
from xml.sax.saxutils import escape
from config import LOG_DIR

# Set up logging configuration
log_file = LOG_DIR / "app.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s - %(message)s",
    handlers=[
        logging.FileHandler(log_file, encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)

logger = logging.getLogger("MultiAgentAssistant")

def get_logger(name: str) -> logging.Logger:
    """Returns a logger child instance for a specific module/agent."""
    return logger.getChild(name)


def export_to_markdown(history: List[Dict[str, Any]]) -> str:
    """Export conversation history as Markdown text."""
    md_lines = ["# Multi-Agent Research Assistant - Conversation Export\n\n"]
    for msg in history:
        role = "User" if msg.get("role") == "user" else "Assistant"
        content = msg.get("content", "")
        route = msg.get("route", "")
        sources = msg.get("sources", [])
        
        md_lines.append(f"### {role}\n")
        if route and role == "Assistant":
            md_lines.append(f"*Execution Route:* `{route}`\n\n")
        md_lines.append(f"{content}\n\n")
        
        if sources and role == "Assistant":
            md_lines.append("**Sources:**\n")
            for src in sources:
                if isinstance(src, dict):
                    title = src.get("title", src.get("source", "Document"))
                    url_or_page = src.get("url", f"Page {src.get('page', 'N/A')}")
                    md_lines.append(f"- [{title}]({url_or_page})\n")
                else:
                    md_lines.append(f"- {src}\n")
            md_lines.append("\n")
        md_lines.append("---\n\n")
    
    return "".join(md_lines)


def export_to_docx(history: List[Dict[str, Any]], filepath: str) -> str:
    """Export conversation history to DOCX format."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        
        doc = Document()
        doc.add_heading("Multi-Agent Research Assistant - Conversation Export", level=0)
        
        for msg in history:
            role = "User" if msg.get("role") == "user" else "Assistant"
            content = msg.get("content", "")
            route = msg.get("route", "")
            sources = msg.get("sources", [])
            
            p_role = doc.add_paragraph()
            run_role = p_role.add_run(f"{role}:")
            run_role.bold = True
            run_role.font.size = Pt(12)
            if role == "User":
                run_role.font.color.rgb = RGBColor(0, 102, 204)
            else:
                run_role.font.color.rgb = RGBColor(0, 153, 76)
                
            if route and role == "Assistant":
                p_route = doc.add_paragraph()
                r_route = p_route.add_run(f"Execution Route: {route}")
                r_route.italic = True
                r_route.font.size = Pt(9)
                r_route.font.color.rgb = RGBColor(128, 128, 128)

            doc.add_paragraph(content)
            
            if sources and role == "Assistant":
                p_src_head = doc.add_paragraph()
                r_src_head = p_src_head.add_run("Sources:")
                r_src_head.bold = True
                for src in sources:
                    p_src = doc.add_paragraph(style='List Bullet')
                    if isinstance(src, dict):
                        title = src.get("title", src.get("source", "Document"))
                        loc = src.get("url", f"Page {src.get('page', 'N/A')}")
                        p_src.add_run(f"{title} ({loc})")
                    else:
                        p_src.add_run(str(src))
            
            doc.add_paragraph("-" * 40)
            
        doc.save(filepath)
        return filepath
    except Exception as e:
        logger.error(f"Error exporting to DOCX: {e}")
        raise e


def export_to_pdf(history: List[Dict[str, Any]], filepath: str) -> str:
    """Export conversation history to PDF format using ReportLab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        
        doc = SimpleDocTemplate(filepath, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=12
        )
        user_style = ParagraphStyle(
            'UserHeader',
            parent=styles['Heading2'],
            fontSize=12,
            textColor=colors.HexColor('#2563EB'),
            spaceBefore=8,
            spaceAfter=4
        )
        assistant_style = ParagraphStyle(
            'AssistantHeader',
            parent=styles['Heading2'],
            fontSize=12,
            textColor=colors.HexColor('#059669'),
            spaceBefore=8,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['BodyText'],
            fontSize=10,
            leading=14,
            spaceAfter=6
        )
        meta_style = ParagraphStyle(
            'Meta',
            parent=styles['Italic'],
            fontSize=9,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=4
        )
        
        story.append(Paragraph("Multi-Agent Research Assistant - Conversation Export", title_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=12))
        
        for msg in history:
            role = "User" if msg.get("role") == "user" else "Assistant"
            content = str(msg.get("content", ""))
            content = re.sub(r"<br\s*/?>", "\n", content, flags=re.IGNORECASE)
            content = escape(content).replace("\n", "<br/>")
            route = escape(str(msg.get("route", "")))
            sources = msg.get("sources", [])
            
            if role == "User":
                story.append(Paragraph("User:", user_style))
            else:
                story.append(Paragraph("Assistant:", assistant_style))
                if route:
                    story.append(Paragraph(f"<i>Execution Route: {route}</i>", meta_style))
            
            story.append(Paragraph(content, body_style))
            
            if sources and role == "Assistant":
                story.append(Paragraph("<b>Sources:</b>", meta_style))
                for src in sources:
                    if isinstance(src, dict):
                        title = escape(str(src.get("title", src.get("source", "Document"))))
                        loc = escape(str(src.get("url", f"Page {src.get('page', 'N/A')}")))
                        story.append(Paragraph(f"• {title} ({loc})", meta_style))
                    else:
                        story.append(Paragraph(f"• {escape(str(src))}", meta_style))
            
            story.append(Spacer(1, 8))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceAfter=8))
            
        doc.build(story)
        return filepath
    except Exception as e:
        logger.error(f"Error exporting to PDF: {e}")
        raise e
